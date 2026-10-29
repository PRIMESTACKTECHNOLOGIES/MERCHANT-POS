import sys, os, json, re, io, base64
from pathlib import Path

DOC_FILES = [
    r"f:\POS ORGINAL\final 1.pdf",
    r"f:\POS ORGINAL\final 2.pptx",
    r"f:\POS ORGINAL\final 3.pptx",
]

results = {}

# ===== PDF extraction with PyMuPDF =====
def extract_pdf_fitz(path):
    import fitz  # PyMuPDF
    doc = fitz.open(path)
    print(f"  Pages: {len(doc)}")
    all_text = ""
    extracted = []
    images = []
    for i, page in enumerate(doc):
        # Extract text
        text = page.get_text("text") or ""
        if text.strip():
            extracted.append({"page": i+1, "text": text})
            all_text += text + "\n\n"
        # Try blocks
        blocks = page.get_text("blocks")
        block_texts = []
        for b in blocks:
            if len(b) >= 5 and b[4].strip():
                block_texts.append(b[4].strip())
        if block_texts and not text.strip():
            bt = "\n".join(block_texts)
            extracted.append({"page": i+1, "method": "blocks", "text": bt})
            all_text += bt + "\n\n"
        # Extract images
        imgs = page.get_images(full=True)
        for j, img in enumerate(imgs):
            try:
                xref = img[0]
                pix = fitz.Pixmap(doc, xref)
                if pix.n >= 5:  # CMYK -> RGB
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                img_bytes = pix.tobytes("png")
                img_path = os.path.join(r"f:\POS ORGINAL\_doc_images", 
                    f"{os.path.basename(path)}_p{i+1}_img{j+1}.png")
                os.makedirs(os.path.dirname(img_path), exist_ok=True)
                with open(img_path, "wb") as f:
                    f.write(img_bytes)
                images.append({"page": i+1, "img_index": j+1, "path": img_path, "size": len(img_bytes)})
                pix = None
            except Exception as e:
                pass
    return extracted, all_text, images

# ===== PPTX extraction with python-pptx, aggressive =====
def extract_pptx_aggressive(path):
    from pptx import Presentation
    from pptx.util import Emu
    import zipfile, xml.etree.ElementTree as ET
    prs = Presentation(path)
    print(f"  Slides: {len(prs.slides)}")
    all_text = ""
    extracted = []
    images = []
    os.makedirs(r"f:\POS ORGINAL\_doc_images", exist_ok=True)

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = []
        # All shapes
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        ptext = "".join([run.text for run in para.runs])
                        if not ptext.strip() and para.text.strip():
                            ptext = para.text
                        if ptext.strip():
                            slide_texts.append(ptext.strip())
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            slide_texts.append(" | ".join(cells))
            except Exception as e:
                pass
        if slide_texts:
            st = "\n".join(slide_texts)
            extracted.append({"slide": slide_num, "text": st})
            all_text += st + "\n\n"

    # Also parse slide XML directly
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    try:
        with zipfile.ZipFile(path, "r") as zf:
            slide_files = sorted([n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)])
            for sfile in slide_files:
                try:
                    slide_num = int(re.search(r"slide(\d+)", sfile).group(1))
                except:
                    slide_num = 0
                with zf.open(sfile) as f:
                    tree = ET.parse(f)
                    root = tree.getroot()
                    xml_texts = []
                    for t in root.iter("{%s}t" % ns["a"]):
                        if t.text and t.text.strip():
                            xml_texts.append(t.text.strip())
                    # Also try p:sp / p:txBody
                    for txbody in root.iter("{%s}txBody" % ns["p"]):
                        for t in txbody.iter("{%s}t" % ns["a"]):
                            if t.text and t.text.strip():
                                pass  # already got above
                    if xml_texts and not any(e.get("slide") == slide_num and e.get("text") for e in extracted):
                        st = "\n".join(xml_texts)
                        already = False
                        for e in extracted:
                            if e.get("slide") == slide_num:
                                already = True
                                if st not in e.get("text", ""):
                                    e["text_xml"] = st
                                break
                        if not already:
                            extracted.append({"slide": slide_num, "method": "xml", "text": st})
                            all_text += st + "\n\n"

            # Extract media (images)
            media_files = [n for n in zf.namelist() if n.startswith("ppt/media/")]
            for mfile in media_files:
                try:
                    img_name = os.path.basename(mfile)
                    img_bytes = zf.read(mfile)
                    img_path = os.path.join(r"f:\POS ORGINAL\_doc_images",
                        f"{os.path.basename(path)}_{img_name}")
                    with open(img_path, "wb") as f:
                        f.write(img_bytes)
                    images.append({"path": img_path, "size": len(img_bytes), "name": img_name})
                except Exception as e:
                    pass
    except Exception as e:
        print(f"  PPTX zip/xml err: {e}")

    return extracted, all_text, images

# ===== Scan for credentials =====
def scan_for_creds(all_text):
    lines = [l.strip() for l in all_text.splitlines() if l.strip()]
    cred_lines = []
    patterns = [
        (re.compile(r"(?i)(api\s*[_-]?\s*key|apikey|secret\s*[_-]?\s*key|access\s*[_-]?\s*key|private\s*[_-]?\s*key|client\s*[_-]?\s*secret)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(password|passwd|pwd|token|jwt|bearer)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(client\s*[_-]?\s*id|merchant\s*[_-]?\s*id|mid|tid|terminal\s*[_-]?\s*id|account\s*[_-]?\s*id)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(username|user\s*[_-]?\s*id|login|email)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(endpoint|url|host|base[_-]?url|gateway|webhook|callback|api\s*[_-]?\s*url)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(hmac|signature|cert|certificate|pem|crt|key\s*file)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(port|sandbox|staging|production|uat|live|test)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(visa|mastercard|unionpay|amex|scheme|network)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"(?i)(swift|iban|routing|sort|ach|wire)\s*[:=]\s*(.+)", re.I), 1, 2),
        (re.compile(r"https?://[^\s\"'<>)]+", re.I), None, None),
        (re.compile(r"\b[A-Fa-f0-9]{32,64}\b"), None, None),
        (re.compile(r"\b\d{13,19}\b"), None, None),
        (re.compile(r"\b(?:pk|sk|rk)_(?:live|test)_[A-Za-z0-9]{10,}\b"), None, None),
        (re.compile(r"\b[A-Za-z0-9_\-\.]{6,}\.(?:pem|crt|key|pfx)\b"), None, None),
    ]
    findings = {}
    for line in lines:
        for pat, gi, vi in patterns:
            for m in pat.finditer(line):
                if gi is None:
                    key = pat.pattern[:30]
                    val = m.group(0)
                else:
                    try:
                        key = m.group(gi).strip()
                        val = m.group(vi).strip() if vi else m.group(0)
                    except:
                        key = pat.pattern[:30]
                        val = m.group(0)
                # Clean val
                val = re.sub(r"^[\"'\s]+|[\"'\s]+$", "", val)
                if len(val) >= 3:
                    bucket = findings.setdefault(key.lower() if len(key) < 60 else "pattern_match", [])
                    if val not in bucket:
                        bucket.append(val)
    return findings

for doc_path in DOC_FILES:
    fname = os.path.basename(doc_path)
    key = fname
    print(f"\n{'='*80}")
    print(f"FILE: {fname}")
    print(f"EXISTS: {os.path.exists(doc_path)}  SIZE: {os.path.getsize(doc_path)} bytes")
    ext = os.path.splitext(fname)[1].lower()
    results[key] = {"type": ext, "extracted": [], "images": [], "credential_findings": {}}

    all_text = ""
    if ext == ".pdf":
        print(f"  Extraction: PyMuPDF")
        try:
            ex, at, imgs = extract_pdf_fitz(doc_path)
            results[key]["extracted"] = ex
            results[key]["images"] = imgs
            all_text = at
            print(f"  Images extracted: {len(imgs)}")
        except Exception as e:
            print(f"  PyMuPDF err: {e}")
    elif ext in (".pptx",):
        print(f"  Extraction: python-pptx aggressive")
        try:
            ex, at, imgs = extract_pptx_aggressive(doc_path)
            results[key]["extracted"] = ex
            results[key]["images"] = imgs
            all_text = at
            print(f"  Images extracted: {len(imgs)}")
        except Exception as e:
            print(f"  PPTX err: {e}")

    print(f"\n--- EXTRACTED TEXT ({key}) ---")
    if all_text.strip():
        print(all_text[:20000])
        if len(all_text) > 20000:
            print(f"\n... [truncated, total {len(all_text)} chars] ...")
    else:
        print("  (no text extracted - document appears to be image-only)")

    # Credential scan
    print(f"\n--- CREDENTIAL SCAN ({key}) ---")
    creds = scan_for_creds(all_text)
    results[key]["credential_findings"] = creds
    if creds:
        for ck, cvals in creds.items():
            print(f"  [{ck}]")
            for v in cvals:
                print(f"    -> {repr(v)}")
    else:
        print("  (no structured credential KV patterns found in text)")

# Save
out = r"f:\POS ORGINAL\_final_docs_extracted_v2.json"
with open(out, "w", encoding="utf-8") as f:
    json.dump(results, f, indent=2, ensure_ascii=False, default=str)
print(f"\nSaved: {out}")
print(f"Images saved to: f:\\POS ORGINAL\\_doc_images")
